"""Phase 7: workpaper export → edit → import (changeset) → apply, and the review packet."""
from __future__ import annotations

import base64
import io

import pytest
from openpyxl import load_workbook

from server import demo, store, workpaper
from server.errors import ToolFailure
from server.tools import load_all

T = load_all()


def call(_tool, *args, **kwargs):
    return T[_tool].fn(*args, **kwargs)


@pytest.fixture(scope="module", autouse=True)
def data_root(tmp_path_factory):
    root = tmp_path_factory.mktemp("data")
    store.configure(root)
    demo.seed()
    yield root
    store._root = None


def approved_case(name="Outputs case"):
    case = call("create_case", name, 2025, "single")["case"]
    doc = call("intake_k1", case["id"], "oak-ventures")["document"]
    for f in call("read_k1", doc["id"])["flags"]:
        if f["ack_required"]:
            call("acknowledge_flag", doc["id"], f["path"])
    call("approve_k1", doc["id"])
    return case, doc


def open_wb(wp):
    return load_workbook(wp["path"])


def k1_sheet(wb):
    return next(wb[s] for s in wb.sheetnames if s.startswith("K-1"))


def row_of(ws, path) -> int:
    return next(r for r in range(2, ws.max_row + 1) if ws.cell(r, workpaper.K1_PATH_COL).value == path)


def set_value(ws, path, value):
    cell = ws.cell(row_of(ws, path), workpaper.K1_VALUE_COL)
    cell.value = value
    return cell.coordinate


def b64(wb) -> str:
    buf = io.BytesIO()
    wb.save(buf)
    return base64.b64encode(buf.getvalue()).decode()


def test_export_structure_and_protection():
    case, doc = approved_case()
    out = call("export_workpaper", case["id"])
    wp = out["workpaper"]
    assert wp["version"] == 1 and wp["url"] == f"/api/outputs/{wp['id']}/download"
    assert out["sources"][0]["ref"] == f"workpaper://{case['id']}/v1"
    wb = open_wb(wp)
    assert wb["_drivkraft"].sheet_state == "hidden"
    assert {"Read me", "1040"} <= set(wb.sheetnames)
    ws = k1_sheet(wb)
    assert ws.protection.sheet
    row = row_of(ws, "part_iii.box_1")
    assert ws.cell(row, workpaper.K1_VALUE_COL).protection.locked is False
    assert ws.cell(row, 1).protection.locked is True
    assert call("export_workpaper", case["id"])["workpaper"]["version"] == 2


def test_round_trip_changes_conflicts_and_invalid():
    case, doc = approved_case("Round trip")
    wp = call("export_workpaper", case["id"])["workpaper"]
    wb = open_wb(wp)
    ws = k1_sheet(wb)
    c1 = set_value(ws, "part_iii.box_1", "(1,250)")            # a loss typed the accountant's way
    set_value(ws, "part_iii.box_5", "abc")                     # not an amount
    set_value(ws, "part_iii.box_2", 888)                       # also changed in the app → conflict
    ws.insert_rows(3)                                          # rows are matched by path, not position
    call("edit_k1_value", doc["id"], "part_iii.box_2", 777, "app edit after export")

    cs = call("import_workpaper", case["id"], b64(wb), "edited.xlsx")["changeset"]
    by = {i["path"]: i for i in cs["items"]}
    assert cs["counts"] == {"changes": 3, "conflicts": 1, "invalid": 1, "applied": 0, "failed": 0, "rejected": 0}
    assert cs["items"][0]["path"] == "part_iii.box_2"          # conflicts pinned first
    assert by["part_iii.box_2"]["conflict"] and by["part_iii.box_2"]["current"] == 777
    assert by["part_iii.box_1"]["new"] == -1250 and not by["part_iii.box_1"]["conflict"]
    assert by["part_iii.box_1"]["cell"] != c1                  # the inserted row moved it
    assert by["part_iii.box_5"]["invalid"]
    # Nothing changed yet (box 1 is still the extracted value; box 2 still the app's edit)
    assert call("get_k1_box", doc["id"], "1")["entries"][0]["value"] != -1250

    with pytest.raises(ToolFailure) as exc:
        call("apply_changeset", cs["id"], [by["part_iii.box_5"]["id"]])
    assert exc.value.code == "invalid_items"

    out = call("apply_changeset", cs["id"], [by["part_iii.box_1"]["id"]])
    assert out["changeset"]["counts"]["applied"] == 1 and out["changeset"]["counts"]["rejected"] == 2
    assert out["k1s"][0]["status"] == "needs_review"          # an approved K-1 goes back to review
    edits = store.list_edits(doc["id"])
    assert edits[-1]["new_value"] == -1250 and "Workpaper v1 import" in edits[-1]["reason"]
    assert call("get_k1_box", doc["id"], "2")["entries"][0]["value"] == 777      # conflict rejected: app wins

    with pytest.raises(ToolFailure) as exc:
        call("apply_changeset", cs["id"], [])
    assert exc.value.code == "changeset_closed"


def test_stale_item_fails_but_others_apply():
    case, doc = approved_case("Stale")
    wb = open_wb(call("export_workpaper", case["id"])["workpaper"])
    ws = k1_sheet(wb)
    set_value(ws, "part_iii.box_1", 111)
    set_value(ws, "part_iii.box_2", 222)
    cs = call("import_workpaper", case["id"], b64(wb))["changeset"]
    call("edit_k1_value", doc["id"], "part_iii.box_2", 999, "edited after the import")
    out = call("apply_changeset", cs["id"], [i["id"] for i in cs["items"]])
    st = {i["path"]: i for i in out["changeset"]["items"]}
    assert st["part_iii.box_1"]["status"] == "applied"
    assert st["part_iii.box_2"]["status"] == "failed" and "after the import" in st["part_iii.box_2"]["error"]


def test_checklist_status_and_discard():
    case, _ = approved_case("Checklist")
    item = store.insert_checklist({"id": store.new_id("chk"), "case_id": case["id"], "item": "Brokerage 1099-B"})
    wb = open_wb(call("export_workpaper", case["id"])["workpaper"])
    ws = wb["Checklist"]
    assert ws.cell(2, workpaper.CHK_ID_COL).value == item["id"]
    ws.cell(2, workpaper.CHK_STATUS_COL).value = "Received"
    cs = call("import_workpaper", case["id"], b64(wb))["changeset"]
    assert [(i["kind"], i["new"]) for i in cs["items"]] == [("checklist_status", "received")]
    call("discard_changeset", cs["id"])
    assert store.get_checklist_item(item["id"])["status"] == "open"
    cs = call("import_workpaper", case["id"], b64(wb))["changeset"]
    call("apply_changeset", cs["id"], ["c1"])
    assert store.get_checklist_item(item["id"])["status"] == "received"


def test_formula_and_unchanged_file():
    case, _ = approved_case("Formula")
    wb = open_wb(call("export_workpaper", case["id"])["workpaper"])
    assert call("import_workpaper", case["id"], b64(wb))["changeset"]["items"] == []
    set_value(k1_sheet(wb), "part_iii.box_1", "=1+1")      # openpyxl saves no cached result
    [item] = call("import_workpaper", case["id"], b64(wb))["changeset"]["items"]
    assert item["invalid"] and "formula" in item["invalid"]


def test_import_refusals():
    case, _ = approved_case("Refusals")
    other, _ = approved_case("Other")
    wb = open_wb(call("export_workpaper", case["id"])["workpaper"])
    for args, code in [((other["id"], b64(wb)), "wrong_case"),
                       ((case["id"], base64.b64encode(b"not a zip").decode()), "not_xlsx"),
                       ((case["id"], "%%%"), "bad_file"),
                       (("ref-k1s", b64(wb)), "case_read_only")]:
        with pytest.raises(ToolFailure) as exc:
            call("import_workpaper", *args)
        assert exc.value.code == code
    from openpyxl import Workbook
    with pytest.raises(ToolFailure) as exc:
        call("import_workpaper", case["id"], b64(Workbook()))
    assert exc.value.code == "not_a_workpaper"


def test_reference_case_exports_and_packet():
    wp = call("export_workpaper", "ref-bench-82")
    assert "read-only" in wp["note"]
    p = call("build_review_packet", "ref-bench-82")["packet"]
    data = open(p["path"], "rb").read()
    assert data.startswith(b"%PDF") and p["pages"] >= 3 and "return" in p["sections"]
    listing = call("list_outputs", "ref-bench-82")
    assert listing["packets"][0]["id"] == p["id"] and listing["workpapers"][0]["kind"] == "workpaper"


def test_packet_without_calculation():
    case = call("create_case", "Empty packet", 2025, "single")["case"]
    p = call("build_review_packet", case["id"])["packet"]
    assert p["pages"] >= 2 and p["filename"] == "empty-packet-review-packet-v1.pdf"


def test_http_download():
    from fastapi.testclient import TestClient
    from server.app import build_http

    with TestClient(build_http()) as client:
        wp = client.post("/api/cases/ref-bench-82/workpapers").json()["workpaper"]
        r = client.get(wp["url"])
        assert r.status_code == 200 and r.content[:2] == b"PK"
        assert "attachment" in r.headers["content-disposition"] and ".xlsx" in r.headers["content-disposition"]
        pk = client.post("/api/cases/ref-bench-82/packets").json()["packet"]
        r = client.get(pk["url"] + "?inline=1")
        assert r.headers["content-type"] == "application/pdf" and r.headers["content-disposition"].startswith("inline")
        assert client.get("/api/outputs/nope/download").json()["error"]["code"] == "output_not_found"
