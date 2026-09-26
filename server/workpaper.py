"""The Excel workpaper (Phase 7): build it from a case snapshot, read an edited copy back.

Two kinds of cell are editable, and only these (the sheets are protected without a
password, so Excel steers people to them):
  K-1 sheets   Value column: one row per OTD path (the path is in the last column)
  Checklist    Status column: open | received
Everything else (the 1040, notes, research, scenarios, follow-ups) is for reading.

A hidden `_drivkraft` sheet records the workpaper id, the case and the value of every
editable cell at export. On import, each cell is compared with that baseline (what
the person changed) and with the case now (what changed since the export). A cell
changed in both places is a conflict.
"""
from __future__ import annotations

import io
import json
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import k1doc
from .errors import ToolFailure

FORMAT = "drivkraft-workpaper/1"
META = "_drivkraft"
MAX_BYTES = 5 * 1024 * 1024

INK, MUTED, LINE = "0F172A", "475569", "CBD5E1"
HEAD = PatternFill("solid", fgColor="F1F5F9")
EDIT = PatternFill("solid", fgColor="FEF9C3")     # the only cells people should type in
WARN = Font(color="B45309", size=10)
TITLE = Font(bold=True, size=14, color=INK)
BOLD = Font(bold=True, color=INK)
SMALL = Font(size=10, color=MUTED)
MONO = Font(name="Menlo", size=10, color=MUTED)
THIN = Border(bottom=Side(style="thin", color=LINE))
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
UNLOCKED = Protection(locked=False)

K1_COLS = (("Box", 12), ("Description", 44), ("Value", 22), ("In calculation", 16), ("OpenTax field", 28),
           ("Flags", 26), ("Original (before edits)", 22), ("OTD path", 30))
K1_VALUE_COL, K1_PATH_COL = 3, 8
CHK_COLS = (("Requested document", 44), ("Detail", 52), ("Status", 14), ("From", 36), ("id", 14))
CHK_STATUS_COL, CHK_ID_COL = 3, 5
DISPOSITION = {"mapped": "Yes", "collapsed": "Yes (combined)", "derived": "Yes (derived)",
               "unsupported": "No", "informational": "Info only"}


def _num_format(v: Any) -> str | None:
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return None
    return "#,##0_);(#,##0);0_)" if float(v).is_integer() else "#,##0.00_);(#,##0.00);0.00_)"


def _sheet_name(base: str, taken: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "", base).strip()[:31] or "Sheet"
    name, n = base, 2
    while name.lower() in taken:
        suffix = f" ({n})"
        name, n = base[:31 - len(suffix)] + suffix, n + 1
    taken.add(name.lower())
    return name


def _header(ws, cols) -> None:
    for i, (title, width) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=title)
        c.font, c.fill, c.border = BOLD, HEAD, THIN
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"


def _protect(ws) -> None:
    ws.protection.sheet = True
    ws.protection.formatColumns = False     # let people widen columns
    ws.protection.formatRows = False
    ws.protection.autoFilter = False


def _table(ws, cols, rows: list[list]) -> None:
    _header(ws, cols)
    for r, row in enumerate(rows, 2):
        for i, v in enumerate(row, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment = WRAP
            if (fmt := _num_format(v)) is not None:
                c.number_format = fmt
    _protect(ws)


def _cell_value(v: Any) -> Any:
    return v if v is None or isinstance(v, (bool, int, float)) else str(v)


# ── Build ─────────────────────────────────────────────────────────────────

def build(snap: dict, workpaper_id: str, version: int) -> tuple[bytes, dict]:
    """The workbook bytes and its baseline (sheet, cell, key → value at export)."""
    case = snap["case"]
    wb = Workbook()
    taken = {META.lower()}
    baseline: list[dict] = []

    ws = wb.active
    ws.title = _sheet_name("Read me", taken)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90
    ws["A1"] = f"{case['name']} · workpaper v{version}"
    ws["A1"].font = TITLE
    info = [("Case", f"{case['name']} ({case['id']})"), ("Tax year", case["tax_year"]),
            ("Filing status", case["filing_status"]), ("Exported", snap["generated"]),
            ("Workpaper", f"v{version} · {workpaper_id}"),
            ("Data", "Synthetic data only. Drivkraft Tax is a practice build, not tax software.")]
    for r, (k, v) in enumerate(info, 3):
        ws.cell(row=r, column=1, value=k).font = SMALL
        ws.cell(row=r, column=2, value=v)
    r = len(info) + 4
    ws.cell(row=r, column=1, value="How to edit").font = BOLD
    steps = [
        "Only the yellow cells can be changed: a K-1's Value column and the Checklist's Status column.",
        "Type amounts as numbers. (300) or -300 is a loss; leave a cell empty to clear it.",
        "Save as .xlsx and import it on the case's Outputs page (or import_workpaper). You'll review every changed "
        "cell before anything is applied, and each applied change is recorded as a K-1 edit with its cell reference.",
        "If the case changed after this export, the import marks those cells as conflicts so nothing is overwritten "
        "without a decision.",
        "Editing an approved K-1 returns it to review.",
    ]
    for i, s in enumerate(steps, r + 1):
        ws.cell(row=i, column=1, value=f"{i - r}.").alignment = TOP
        ws.cell(row=i, column=2, value=s).alignment = WRAP
    r += len(steps) + 2
    ws.cell(row=r, column=1, value="Sheets").font = BOLD

    sheet_rows: list[tuple[str, str]] = []

    for k1 in snap["k1s"]:
        ks = wb.create_sheet(_sheet_name(f"K-1 {k1['name']}", taken))
        _header(ks, K1_COLS)
        for row, e in enumerate(k1["entries"], 2):
            val = _cell_value(e["value"])
            cells = [e["label"], e["description"], val, DISPOSITION.get(e["disposition"], e["disposition"]),
                     e["field"], ", ".join(e["flags"]) or None, _cell_value(e["original"]) if e["edited"] else None,
                     e["path"]]
            for col, v in enumerate(cells, 1):
                c = ks.cell(row=row, column=col, value=v)
                c.alignment = WRAP if col in (2, 3) else TOP
                if (fmt := _num_format(v)) is not None:
                    c.number_format = fmt
            vc = ks.cell(row=row, column=K1_VALUE_COL)
            vc.fill, vc.protection = EDIT, UNLOCKED
            ks.cell(row=row, column=K1_PATH_COL).font = MONO
            if e["flags"]:
                ks.cell(row=row, column=6).font = WARN
            baseline.append({"sheet": ks.title, "cell": vc.coordinate, "kind": "k1_value", "doc_id": k1["id"],
                             "key": e["path"], "value": e["value"]})
        _protect(ks)
        status = {"approved": "approved: an edit here returns it to review", "needs_review": "in review",
                  "blocked": "blocked by bridge errors"}.get(k1["status"], k1["status"])
        sheet_rows.append((ks.title, f"{k1['name']}: {len(k1['entries'])} values, {status}."))

    if snap["checklist"]:
        cs = wb.create_sheet(_sheet_name("Checklist", taken))
        _header(cs, CHK_COLS)
        dv = DataValidation(type="list", formula1='"open,received"', allow_blank=False, showErrorMessage=True,
                            errorTitle="Status", error="Use open or received.")
        cs.add_data_validation(dv)
        for row, c in enumerate(snap["checklist"], 2):
            for col, v in enumerate((c["item"], c["detail"], c["status"], c["from"], c["id"]), 1):
                cs.cell(row=row, column=col, value=v).alignment = WRAP
            sc = cs.cell(row=row, column=CHK_STATUS_COL)
            sc.fill, sc.protection = EDIT, UNLOCKED
            dv.add(sc)
            cs.cell(row=row, column=CHK_ID_COL).font = MONO
            baseline.append({"sheet": cs.title, "cell": sc.coordinate, "kind": "checklist_status",
                             "key": c["id"], "value": c["status"]})
        _protect(cs)
        sheet_rows.append((cs.title, "Documents requested in meetings. Status is editable (open or received)."))

    ret = snap["return"]
    fs = wb.create_sheet(_sheet_name("1040", taken))
    if ret:
        _table(fs, (("Line", 8), ("Description", 44), ("Amount", 18), ("Section", 16)),
               [[l["line"], l["label"], l["value"], l["section"]] for l in ret["lines"]])
        n = len(ret["lines"]) + 3
        for i, cv in enumerate(ret["caveats"]):
            fs.cell(row=n + i, column=2, value=cv["message"]).font = WARN
        sheet_rows.append((fs.title, f"The OpenTax 1040 for the approved K-1s: {len(ret['lines'])} lines. Read-only."))
    else:
        fs["A1"] = f"No calculation yet: {snap['return_error'] or 'nothing approved'}."
        _protect(fs)
        sheet_rows.append((fs.title, "The 1040 (empty until a K-1 is approved)."))

    if snap["scenarios"]:
        ss = wb.create_sheet(_sheet_name("Scenarios", taken))
        _table(ss, (("Scenario", 28), ("Changes", 80), ("Saved", 22)),
               [[s["name"], json.dumps(s["changes"]), s["created"]] for s in snap["scenarios"]])
        sheet_rows.append((ss.title, "Saved what-ifs. Re-run them on the Return page."))

    if snap["notes"]:
        ns = wb.create_sheet(_sheet_name("Notes", taken))
        rows = []
        for n in snap["notes"]:
            decisions = "\n".join(f"{d['at'] + ' ' if d['at'] else ''}{d['text']}" for d in n["decisions"])
            rows.append([n["meeting_date"], n["title"], ", ".join(n["attendees"]),
                         n["summary"] or ("Not analyzed yet" if not n["analyzed"] else None), decisions or None])
        _table(ns, (("Date", 12), ("Meeting", 32), ("Attendees", 28), ("Summary", 60), ("Decisions", 60)), rows)
        sheet_rows.append((ns.title, "Meeting notes: summaries and decisions with timestamps."))

    if snap["research"]:
        rs = wb.create_sheet(_sheet_name("Research", taken))
        rows = []
        for q in snap["research"]:
            cites = "\n".join(f"[{i}] {c['label']}" + (f" · {c['url']}" if c.get("url") else "")
                              for i, c in enumerate(q["citations"], 1))
            rows.append([q["question"], q["mode"], "Cached (written for this demo)" if q["cached"] else "Live Bizora",
                         q["cost_usd"], cites])
        _table(rs, (("Question", 50), ("Mode", 8), ("Source", 28), ("Cost (USD)", 11), ("Citations", 80)), rows)
        for r in range(2, len(rows) + 2):
            rs.cell(row=r, column=4).number_format = "$0.00"
        sheet_rows.append((rs.title, "Tax research questions, with citations. Full answers are in the review packet."))

    if snap["follow_ups"]:
        us = wb.create_sheet(_sheet_name("Follow-ups", taken))
        _table(us, (("Subject", 40), ("Email", 90), ("From meeting", 30), ("Approved", 22)),
               [[f["subject"], f["body"], f["note"], f["approved"]] for f in snap["follow_ups"]])
        sheet_rows.append((us.title, "Follow-up emails approved in the Inbox."))

    for i, (name, desc) in enumerate(sheet_rows, r + 1):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=desc).alignment = WRAP
    _protect(ws)

    meta = wb.create_sheet(META)
    meta.sheet_state = "hidden"
    head = {"format": FORMAT, "workpaper_id": workpaper_id, "case_id": case["id"], "version": version,
            "exported": snap["generated"]}
    for i, (k, v) in enumerate(head.items(), 1):
        meta.cell(row=i, column=1, value=k)
        meta.cell(row=i, column=2, value=v)
    start = len(head) + 2
    for i, h in enumerate(("sheet", "cell", "kind", "doc_id", "key", "value"), 1):
        meta.cell(row=start, column=i, value=h)
    for r, b in enumerate(baseline, start + 1):
        for i, v in enumerate((b["sheet"], b["cell"], b["kind"], b.get("doc_id"), b["key"], json.dumps(b["value"])), 1):
            meta.cell(row=r, column=i, value=v)
    _protect(meta)

    buf = io.BytesIO()
    wb.save(buf)
    sheets = [t for t in wb.sheetnames if t != META]
    return buf.getvalue(), {"sheets": sheets, "editable_cells": len(baseline)}


# ── Read back ─────────────────────────────────────────────────────────────

def read(data: bytes) -> dict:
    """The workbook's header, baseline, and the values now in its editable cells.

    Rows are found by their key (the OTD path, the checklist id), not by position,
    so sorting or inserting rows doesn't misread a cell."""
    if len(data) > MAX_BYTES:
        raise ToolFailure("file_too_large", "Workpapers are limited to 5 MB", "Import the workbook this app exported.",
                          status=413)
    try:
        wb = load_workbook(io.BytesIO(data))
    except Exception as exc:
        raise ToolFailure("not_xlsx", "That file isn't an Excel workbook (.xlsx)",
                          "Save it as .xlsx (Excel Workbook) and import again.", status=422) from exc
    if META not in wb.sheetnames:
        raise ToolFailure("not_a_workpaper", "This workbook wasn't exported by Drivkraft Tax",
                          "Export a workpaper from the case, edit that file, and import it.", status=422)
    meta = wb[META]
    head: dict[str, Any] = {}
    row = 1
    while (k := meta.cell(row=row, column=1).value) not in (None, "sheet"):
        head[k] = meta.cell(row=row, column=2).value
        row += 1
    if head.get("format") != FORMAT:
        raise ToolFailure("unsupported_format", f"Workpaper format {head.get('format')!r} isn't supported",
                          "Export a fresh workpaper and edit that one.", status=422)
    while meta.cell(row=row, column=1).value != "sheet":
        row += 1
    baseline = []
    for r in meta.iter_rows(min_row=row + 1, values_only=True):
        if not r or r[0] is None:
            continue
        sheet, cell, kind, doc_id, key, value = r[:6]
        baseline.append({"sheet": sheet, "cell": cell, "kind": kind, "doc_id": doc_id, "key": key,
                         "value": json.loads(value) if value is not None else None})

    computed = None     # data_only copy, loaded only if someone typed a formula

    def value_at(ws, r: int, c: int) -> tuple[Any, str | None]:
        nonlocal computed
        v = ws.cell(row=r, column=c).value
        if isinstance(v, str) and v.startswith("="):
            computed = computed or load_workbook(io.BytesIO(data), data_only=True)
            v = computed[ws.title].cell(row=r, column=c).value
            if v is None:
                return None, "The cell has a formula with no saved result. Type the value instead."
        if isinstance(v, str) and not v.strip():
            v = None
        return v, None

    cells: dict[tuple[str, str], dict] = {}
    missing_sheets = sorted({b["sheet"] for b in baseline if b["sheet"] not in wb.sheetnames})
    for name in {b["sheet"] for b in baseline} - set(missing_sheets):
        ws = wb[name]
        kinds = {b["kind"] for b in baseline if b["sheet"] == name}
        key_col, val_col = (K1_PATH_COL, K1_VALUE_COL) if "k1_value" in kinds else (CHK_ID_COL, CHK_STATUS_COL)
        for r in range(2, ws.max_row + 1):
            key = ws.cell(row=r, column=key_col).value
            if not isinstance(key, str) or not key.strip():
                continue
            v, err = value_at(ws, r, val_col)
            cells[(name, key.strip())] = {"cell": ws.cell(row=r, column=val_col).coordinate, "value": v, "error": err}
    return {"head": head, "baseline": baseline, "cells": cells, "missing_sheets": missing_sheets}


def same(a: Any, b: Any) -> bool:
    """Equal as K-1 values: 1250 == 1250.0; text compared after trimming."""
    if isinstance(a, bool) or isinstance(b, bool):
        return a is b or (isinstance(a, bool) and isinstance(b, bool) and a == b)
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(float(a) - float(b)) < 0.005
    if isinstance(a, str) and isinstance(b, str):
        return a.strip() == b.strip()
    return a == b


def parse_k1(raw: Any, baseline: Any) -> Any:
    """A typed cell → the K-1 value, the way edit_k1_value would read it."""
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    return k1doc.coerce(raw, baseline)
